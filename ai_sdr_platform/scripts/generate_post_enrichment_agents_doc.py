from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = Path("ai_sdr_platform/docs/post_enrichment_agents.docx")


AGENTS = [
    {
        "name": "Intent Agent",
        "primary_outcome": "Detect which discovered and enriched accounts are showing real buying intent right now.",
        "inputs": [
            "Enriched account brief",
            "Website activity",
            "Email or campaign engagement",
            "Product-interest events",
            "News, hiring, or funding signals",
        ],
        "outputs": [
            "Intent score",
            "Intent reasons",
            "Urgency tier such as hot, warm, or cold",
            "Recommended timing for outreach",
        ],
        "owner": "Shared between Marketing and Sales",
        "technical": "Combines event signals with enriched account context, scores urgency, and writes a structured intent object for downstream ranking.",
        "marketing_related": "High",
    },
    {
        "name": "Lead Scoring Agent",
        "primary_outcome": "Rank accounts and contacts so reps and automation focus on the best opportunities first.",
        "inputs": [
            "ICP fit",
            "Discovery fit score",
            "Enrichment quality",
            "Intent score",
            "Persona match",
        ],
        "outputs": [
            "Lead score",
            "Priority band",
            "Scoring explanation",
            "Queue ordering for outreach",
        ],
        "owner": "Shared between Marketing and Sales",
        "technical": "Calculates a weighted score from structured upstream outputs and stores both the final score and the reasoning trail.",
        "marketing_related": "High",
    },
    {
        "name": "Personalization Agent",
        "primary_outcome": "Turn research into sharp account-specific messaging angles for each persona.",
        "inputs": [
            "Enrichment brief",
            "Intent reasons",
            "Persona data",
            "Offer summary",
            "Past campaign context",
        ],
        "outputs": [
            "Email angle",
            "LinkedIn angle",
            "Call opener",
            "Persona-specific value proposition",
        ],
        "owner": "Shared, usually Marketing-enabled and Sales-used",
        "technical": "Uses LLM synthesis plus structured context to generate short reusable message blocks, not full campaign execution.",
        "marketing_related": "High",
    },
    {
        "name": "Outreach Agent",
        "primary_outcome": "Convert messaging into channel-ready outbound actions.",
        "inputs": [
            "Personalization package",
            "Preferred channels",
            "Contact details",
            "Compliance or suppression rules",
        ],
        "outputs": [
            "Prepared email",
            "Prepared LinkedIn message",
            "SMS or WhatsApp copy where allowed",
            "Outbound action record",
        ],
        "owner": "Sales-owned with Marketing guardrails",
        "technical": "Maps a message package to the chosen channel, applies templates and constraints, and hands off to execution systems.",
        "marketing_related": "Medium to High",
    },
    {
        "name": "Follow-Up Agent",
        "primary_outcome": "Manage sequence timing, retries, and stop conditions after first outreach.",
        "inputs": [
            "Outreach history",
            "Reply status",
            "Intent tier",
            "Sequence rules",
        ],
        "outputs": [
            "Next follow-up step",
            "Delay interval",
            "Channel switch recommendation",
            "Stop or continue decision",
        ],
        "owner": "Sales Operations / Revenue Operations",
        "technical": "Implements deterministic sequence policy with optional AI adjustments based on recent signals and responses.",
        "marketing_related": "Medium",
    },
    {
        "name": "Conversation Agent",
        "primary_outcome": "Understand replies and route each response to the right next action.",
        "inputs": [
            "Inbound email or message reply",
            "Prior outreach context",
            "Account and persona brief",
        ],
        "outputs": [
            "Reply classification",
            "Objection or interest label",
            "Suggested response type",
            "Escalation path if needed",
        ],
        "owner": "Sales-owned",
        "technical": "Classifies inbound text such as interested, not now, pricing request, objection, or unsubscribe and writes structured next-step metadata.",
        "marketing_related": "Low to Medium",
    },
    {
        "name": "Qualification Agent",
        "primary_outcome": "Decide whether the opportunity is sales-ready and worth rep time.",
        "inputs": [
            "Conversation outcome",
            "ICP fit",
            "Intent score",
            "Account and persona data",
            "Qualification framework such as BANT or MEDDIC",
        ],
        "outputs": [
            "Qualified or not qualified decision",
            "Qualification score",
            "Missing qualification fields",
            "Recommended handoff path",
        ],
        "owner": "Sales-owned",
        "technical": "Runs a rules-plus-LLM evaluation against a chosen qualification framework and records a structured qualification decision.",
        "marketing_related": "Low",
    },
    {
        "name": "Meeting Agent",
        "primary_outcome": "Book meetings and package the context for the account executive or SDR handoff.",
        "inputs": [
            "Qualified lead decision",
            "Calendar availability",
            "Prospect brief",
            "Conversation context",
        ],
        "outputs": [
            "Meeting booking",
            "Agenda",
            "Prospect brief",
            "Owner handoff record",
        ],
        "owner": "Sales-owned",
        "technical": "Coordinates scheduling, creates the meeting object, and assembles the brief needed by the rep taking the call.",
        "marketing_related": "Low",
    },
    {
        "name": "CRM Agent",
        "primary_outcome": "Write all important SDR activity and state changes back into core systems.",
        "inputs": [
            "Discovery, enrichment, scoring, outreach, and meeting outputs",
            "CRM object identifiers",
        ],
        "outputs": [
            "Lead, contact, and account updates",
            "Activity timeline records",
            "Opportunity state updates",
            "Attribution-ready fields",
        ],
        "owner": "Revenue Operations / CRM Operations",
        "technical": "Persists structured agent outputs into Salesforce, HubSpot, or another CRM system while preserving auditability.",
        "marketing_related": "Medium to High",
    },
    {
        "name": "Analytics Agent",
        "primary_outcome": "Measure what is working across the SDR funnel and explain where to improve.",
        "inputs": [
            "Outreach events",
            "Reply events",
            "Meeting events",
            "CRM outcomes",
            "Campaign and intent signals",
        ],
        "outputs": [
            "Reply rate",
            "Meeting rate",
            "Pipeline contribution",
            "Segment and channel insights",
            "Agent quality metrics",
        ],
        "owner": "Shared between Marketing, Sales, and RevOps",
        "technical": "Aggregates structured events from the full workflow and turns them into dashboards, summaries, and optimization recommendations.",
        "marketing_related": "High",
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_page_layout(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


def style_base(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color in [
        ("Heading 1", 16, RGBColor(0x2E, 0x74, 0xB5)),
        ("Heading 2", 13, RGBColor(0x2E, 0x74, 0xB5)),
        ("Heading 3", 12, RGBColor(0x1F, 0x4D, 0x78)),
    ]:
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color


def add_title_block(document: Document) -> None:
    title = document.add_paragraph()
    title.style = document.styles["Normal"]
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("AI SDR Platform: Agents After Enrichment")
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)
    title.paragraph_format.space_after = Pt(4)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle_run = subtitle.add_run(
        f"Meeting brief | Generated on {date.today().isoformat()} | Scope: post-Enrichment SDR workflow"
    )
    subtitle_run.font.name = "Calibri"
    subtitle_run.font.size = Pt(10)
    subtitle_run.italic = True
    subtitle_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    subtitle.paragraph_format.space_after = Pt(12)


def add_lead_summary(document: Document) -> None:
    document.add_heading("Where We Are Now", level=1)
    p = document.add_paragraph()
    p.add_run(
        "The first three agents already in progress are ICP Agent, Prospect Discovery Agent, and Enrichment Agent. "
        "After Enrichment, the workflow shifts from research into prioritization, message creation, execution, qualification, "
        "CRM synchronization, and performance measurement."
    )
    p = document.add_paragraph()
    p.add_run("Recommended next build order: ").bold = True
    p.add_run(
        "Intent Agent, Lead Scoring Agent, Personalization Agent, Outreach Agent, Follow-Up Agent, "
        "Conversation Agent, Qualification Agent, Meeting Agent, CRM Agent, Analytics Agent."
    )


def add_summary_table(document: Document) -> None:
    document.add_heading("Post-Enrichment Agent Map", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.8), Inches(2.3), Inches(1.55), Inches(0.85)]
    headers = ["Agent", "Primary Outcome", "Owner / Team", "Marketing"]
    for idx, (cell, label) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = widths[idx]
        set_cell_shading(cell, "F2F4F7")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(10)

    for agent in AGENTS:
        row = table.add_row().cells
        values = [
            agent["name"],
            agent["primary_outcome"],
            agent["owner"],
            agent["marketing_related"],
        ]
        for idx, value in enumerate(values):
            row[idx].width = widths[idx]
            p = row[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            r.font.size = Pt(10)


def add_agent_sections(document: Document) -> None:
    document.add_heading("What Each Agent Does", level=1)
    for agent in AGENTS:
        document.add_heading(agent["name"], level=2)

        lead = document.add_paragraph()
        lead.add_run("Primary outcome: ").bold = True
        lead.add_run(agent["primary_outcome"])

        tech = document.add_paragraph()
        tech.add_run("Technical role: ").bold = True
        tech.add_run(agent["technical"])

        document.add_paragraph("Typical inputs:", style="Heading 3")
        for item in agent["inputs"]:
            document.add_paragraph(item, style="List Bullet")

        document.add_paragraph("Typical outputs:", style="Heading 3")
        for item in agent["outputs"]:
            document.add_paragraph(item, style="List Bullet")

        owner = document.add_paragraph()
        owner.add_run("Owner / team: ").bold = True
        owner.add_run(agent["owner"])

        marketing = document.add_paragraph()
        marketing.add_run("Marketing relationship: ").bold = True
        marketing.add_run(agent["marketing_related"])


def add_marketing_section(document: Document) -> None:
    document.add_heading("Agents Most Related to the Marketing Team", level=1)
    intro = document.add_paragraph()
    intro.add_run(
        "These are the agents most likely to be shared with, informed by, or partly owned by the marketing team."
    )

    marketing_agents = [
        ("Intent Agent", "Consumes campaign engagement, website behavior, content interaction, and demand signals."),
        ("Lead Scoring Agent", "Shapes how marketing-qualified activity becomes sales-prioritized work."),
        ("Personalization Agent", "Depends on messaging, positioning, and segment-level context often created by marketing."),
        ("Outreach Agent", "May reuse campaign templates, brand-safe wording, and multi-channel playbooks."),
        ("CRM Agent", "Often writes campaign source, lifecycle, attribution, and suppression data used by marketing."),
        ("Analytics Agent", "Measures conversion, response, and pipeline impact across campaigns and segments."),
    ]
    for name, detail in marketing_agents:
        p = document.add_paragraph(style="List Bullet")
        p.add_run(f"{name}: ").bold = True
        p.add_run(detail)


def add_recommended_order(document: Document) -> None:
    document.add_heading("Recommended Build Order After Enrichment", level=1)
    ordered = [
        "Intent Agent",
        "Lead Scoring Agent",
        "Personalization Agent",
        "Outreach Agent",
        "Follow-Up Agent",
        "Conversation Agent",
        "Qualification Agent",
        "Meeting Agent",
        "CRM Agent",
        "Analytics Agent",
    ]
    for item in ordered:
        document.add_paragraph(item, style="List Number")

    note = document.add_paragraph()
    note.add_run("Why this order works: ").bold = True
    note.add_run(
        "It moves from account understanding into prioritization first, then message generation, then execution, "
        "then response handling, then CRM and analytics. That keeps the system operationally coherent."
    )


def build_doc() -> Path:
    document = Document()
    set_page_layout(document)
    style_base(document)
    add_title_block(document)
    add_lead_summary(document)
    add_summary_table(document)
    add_agent_sections(document)
    add_marketing_section(document)
    add_recommended_order(document)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_doc()
    print(path)
